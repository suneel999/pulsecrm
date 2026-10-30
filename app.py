from __future__ import annotations

import hashlib
import hmac
import io
import logging
import os
import time
from functools import wraps
from pathlib import Path

from dotenv import load_dotenv
from flask import (
    Flask,
    Response,
    abort,
    flash,
    redirect,
    render_template,
    request,
    send_file,
    session,
    url_for,
)
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from werkzeug.middleware.proxy_fix import ProxyFix

load_dotenv()

import app_api
import db
import wa_flow
import wa_flow_crypto

log = logging.getLogger("pulse-desk")

app = Flask(__name__)
app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)
app.secret_key = os.getenv("FLASK_SECRET", "pulse-desk-dev-secret")
app.config["PREFERRED_URL_SCHEME"] = "https"

_last_app_sync = 0.0
_SYNC_GAP_SEC = 30
_purged_old_app = False


def env(name: str, default: str = "") -> str:
    return (os.getenv(name) or default).strip()


def login_required(view):
    @wraps(view)
    def wrapped(*args, **kwargs):
        if not session.get("staff"):
            return redirect(url_for("login", next=request.path))
        return view(*args, **kwargs)

    return wrapped


@app.context_processor
def inject():
    return {
        "STATUSES": db.STATUSES,
        "SERVICES": db.SERVICES,
        "format_when": db.format_when,
        "format_slot": db.format_slot,
        "format_day_label": db.format_day_label,
        "short_doctor": db.short_doctor,
        "source_label": source_label,
        "today": db.today(),
    }


def source_label(source: str) -> str:
    key = (source or "").lower()
    if key == "app":
        return "App"
    if key == "website":
        return "Web"
    if key == "walk-in":
        return "Walk-in"
    if key == "whatsapp":
        return "WA"
    if key == "phone":
        return "Phone"
    return source or "—"


def sync_app_requests(force: bool = False) -> dict | None:
    global _last_app_sync, _purged_old_app
    if not env("APP_API_EMAIL") or not env("APP_API_PASSWORD"):
        return None
    now = time.time()
    if not force and _last_app_sync and (now - _last_app_sync) < _SYNC_GAP_SEC:
        return {"skipped": True, "stats": db.stats()}
    cutoff = app_api.sync_from_date()
    if not _purged_old_app:
        purged = db.purge_app_before(cutoff)
        _purged_old_app = True
    else:
        purged = 0
    rows = app_api.fetch_appointments(limit=50, max_pages=20)
    result = db.upsert_app(rows)
    _last_app_sync = now
    result["fetched"] = len(rows)
    result["purged"] = purged
    result["from"] = cutoff
    result["stats"] = db.stats()
    return result


@app.get("/health")
def health():
    return {"ok": True, "service": "pulse-desk"}


@app.post("/api/ingest")
def ingest():
    key = request.headers.get("X-CRM-Key", "")
    expected = env("CRM_INGEST_KEY")
    if not expected or key != expected:
        abort(401)
    payload = request.get_json(silent=True) or {}
    try:
        row = db.create(payload)
    except ValueError as exc:
        return {"ok": False, "error": str(exc)}, 400
    return {"ok": True, "id": row["id"]}


def _wa_private_key() -> str:
    pem = env("WA_FLOW_PRIVATE_KEY")
    if pem:
        return pem.replace("\\n", "\n")
    path = env("WA_FLOW_PRIVATE_KEY_PATH", "secrets/wa_flow_private.pem")
    key_path = Path(path)
    if not key_path.is_absolute():
        key_path = Path(__file__).resolve().parent / key_path
    if key_path.is_file():
        return key_path.read_text(encoding="utf-8")
    return ""


def _wa_signature_ok(raw_body: bytes) -> bool:
    app_secret = env("WA_APP_SECRET") or env("META_APP_SECRET")
    if not app_secret:
        # Allow first-time wiring without signature; set WA_APP_SECRET before production.
        log.warning("WA_APP_SECRET not set — skipping signature check")
        return True
    header = request.headers.get("X-Hub-Signature-256", "")
    if not header.startswith("sha256="):
        return False
    expected = hmac.new(app_secret.encode("utf-8"), raw_body, hashlib.sha256).hexdigest()
    received = header.removeprefix("sha256=")
    return hmac.compare_digest(expected, received)


@app.get("/api/whatsapp/flow")
def whatsapp_flow_info():
    return {
        "ok": True,
        "service": "pulse-whatsapp-flow",
        "endpoint": "POST /api/whatsapp/flow",
        "hint": "Configure this URL in WhatsApp Manager → Flow → Endpoint",
    }


@app.post("/api/whatsapp/flow")
def whatsapp_flow():
    """Meta WhatsApp Flows data endpoint (encrypted). Saves bookings to desk as source=whatsapp."""
    raw = request.get_data()
    if not _wa_signature_ok(raw):
        return Response(status=432)

    private_pem = _wa_private_key()
    if not private_pem:
        log.error("WA Flow private key missing")
        return Response(status=500)

    body = request.get_json(silent=True) or {}
    try:
        decrypted = wa_flow_crypto.decrypt_request(
            body, private_pem, env("WA_FLOW_PASSPHRASE")
        )
    except wa_flow_crypto.FlowCryptoError as exc:
        log.warning("WA Flow crypto error: %s", exc)
        return Response(status=exc.status_code)

    decrypted_body = decrypted["decrypted_body"]
    try:
        screen_response = wa_flow.get_next_screen(decrypted_body)
    except Exception:
        log.exception("WA Flow logic error body=%s", decrypted_body)
        return Response(status=500)

    encrypted = wa_flow_crypto.encrypt_response(
        screen_response, decrypted["aes_key"], decrypted["initial_vector"]
    )
    return Response(encrypted, mimetype="text/plain")


@app.route("/login", methods=["GET", "POST"])
def login():
    if session.get("staff"):
        return redirect(url_for("desk"))
    error = ""
    if request.method == "POST":
        user = request.form.get("user", "")
        password = request.form.get("pass", "")
        admin_pass = env("ADMIN_PASS")
        if admin_pass and user == env("ADMIN_USER", "pulse") and password == admin_pass:
            session["staff"] = user
            return redirect(request.args.get("next") or url_for("desk"))
        error = "Wrong username or password."
    return render_template("login.html", error=error)


@app.get("/logout")
def logout():
    session.clear()
    return redirect(url_for("login"))


def _desk_query():
    view = request.args.get("view") or "fresh"
    if view not in db.VIEWS:
        view = "fresh"
    try:
        page = max(1, int(request.args.get("page") or 1))
    except ValueError:
        page = 1
    day = (request.args.get("day") or "").strip()
    if not db.valid_day(day):
        # Fresh / coming default to today so staff can flip to previous days
        day = db.today() if view in ("fresh", "coming") else ""
    return {
        "view": view,
        "status": request.args.get("status") or "",
        "q": (request.args.get("q") or "").strip(),
        "day": day,
        "page": page,
    }


def _desk_url(**overrides):
    params = _desk_query()
    params.update(overrides)
    params = {key: value for key, value in params.items() if value not in ("", None)}
    return url_for("desk", **params)


@app.get("/")
@login_required
def desk():
    sync_error = ""
    try:
        sync_app_requests(force=False)
    except Exception as exc:
        sync_error = str(exc)
    query = _desk_query()
    rows, total = db.list_page(
        view=query["view"],
        status=query["status"],
        q=query["q"],
        page=query["page"],
        day=query["day"],
    )
    pages = max(1, (total + db.PAGE_SIZE - 1) // db.PAGE_SIZE)
    if query["page"] > pages:
        query["page"] = pages
        rows, total = db.list_page(
            view=query["view"],
            status=query["status"],
            q=query["q"],
            page=pages,
            day=query["day"],
        )
    return render_template(
        "desk.html",
        rows=rows,
        stats=db.stats(query["day"]),
        total=total,
        pages=pages,
        page=query["page"],
        view=query["view"],
        status=query["status"],
        q=query["q"],
        day=query["day"],
        day_label=db.format_day_label(query["day"] or db.today()),
        prev_day=db.shift_day(query["day"] or db.today(), -1),
        next_day=db.shift_day(query["day"] or db.today(), 1),
        month=db.today()[:7],
        next_url=_desk_url(),
        page_url=_desk_url,
        sync_error=sync_error,
        sync_from=app_api.sync_from_date(),
    )


@app.post("/sync-app")
@login_required
def sync_app():
    try:
        result = sync_app_requests(force=True) or {}
        if result.get("skipped"):
            flash("App list is already up to date.")
        else:
            extra = ""
            if result.get("purged"):
                extra = f" Cleared {result['purged']} old App rows before {result.get('from')}."
            flash(
                f"App sync done: {result.get('added', 0)} new, "
                f"{result.get('refreshed', 0)} refreshed "
                f"(from {result.get('from')}).{extra}"
            )
    except Exception as exc:
        flash(f"App sync failed: {exc}")
    return redirect(request.form.get("next") or url_for("desk", view="fresh"))


@app.get("/api/sync-app")
@login_required
def sync_app_api():
    """Silent sync for the open desk tab — front office does not need to click Sync."""
    try:
        result = sync_app_requests(force=bool(request.args.get("force"))) or {}
        return {
            "ok": True,
            "skipped": bool(result.get("skipped")),
            "added": int(result.get("added") or 0),
            "refreshed": int(result.get("refreshed") or 0),
            "purged": int(result.get("purged") or 0),
            "fetched": int(result.get("fetched") or 0),
            "from": result.get("from") or app_api.sync_from_date(),
            "stats": result.get("stats") or db.stats(),
        }
    except Exception as exc:
        return {"ok": False, "error": str(exc)}, 500


@app.post("/<int:appt_id>/update")
@login_required
def update(appt_id: int):
    payload = {"status": request.form.get("status")}
    if "notes" in request.form:
        payload["notes"] = request.form.get("notes")
    if request.form.get("save_details"):
        payload["services"] = request.form.getlist("services")
        payload["service_other"] = request.form.get("service_other")
    try:
        db.update(appt_id, payload)
    except ValueError as exc:
        flash(str(exc))
    return redirect(request.form.get("next") or url_for("desk"))


@app.post("/add")
@login_required
def add_manual():
    try:
        db.create(
            {
                "name": request.form.get("name"),
                "phone": request.form.get("phone"),
                "email": request.form.get("email"),
                "department": request.form.get("department"),
                "doctor": request.form.get("doctor"),
                "preferred_date": request.form.get("preferred_date"),
                "preferred_time": request.form.get("preferred_time"),
                "notes": request.form.get("notes"),
                "source": request.form.get("source") or "phone",
                "page": "front-office",
            }
        )
        flash("Appointment added.")
    except ValueError as exc:
        flash(str(exc))
    return redirect(url_for("desk", view="fresh"))


@app.get("/report.xlsx")
@login_required
def report():
    month = (request.args.get("month") or db.today()[:7]).strip()
    rows = db.list_appts(month=month)
    book = Workbook()
    sheet = book.active
    sheet.title = month
    headers = [
        "Received",
        "Name",
        "Mobile",
        "Email",
        "Department",
        "Doctor",
        "Preferred date",
        "Preferred time",
        "Status",
        "Services",
        "Other service",
        "Notes",
        "Source",
        "App ID",
    ]
    sheet.append(headers)
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill("solid", fgColor="012D3A")
    thin = Border(
        left=Side(style="thin", color="E5E7EB"),
        right=Side(style="thin", color="E5E7EB"),
        top=Side(style="thin", color="E5E7EB"),
        bottom=Side(style="thin", color="E5E7EB"),
    )
    for cell in sheet[1]:
        cell.font = header_font
        cell.fill = header_fill
    for row in rows:
        sheet.append(
            [
                row["created_at"],
                row["name"],
                row["phone"],
                row["email"],
                row["department"],
                row["doctor"],
                row["preferred_date"],
                row["preferred_time"],
                dict(db.STATUSES).get(row["status"], row["status"]),
                ", ".join(row.get("services_list") or []),
                row["service_other"],
                row["notes"],
                row["source"],
                row.get("external_id") or "",
            ]
        )
    visited = sum(1 for row in rows if row["status"] == "visited")
    sheet.append([])
    sheet.append(["Total", len(rows), "", "", "", "", "", "", f"Visited {visited}", "", "", "", "", ""])
    widths = [20, 28, 16, 28, 22, 36, 14, 14, 12, 28, 18, 36, 12, 12]
    for col, width in enumerate(widths, 1):
        sheet.column_dimensions[sheet.cell(1, col).column_letter].width = width
    for row in sheet.iter_rows(min_row=1, max_row=sheet.max_row, max_col=14):
        for cell in row:
            cell.border = thin
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    buf = io.BytesIO()
    book.save(buf)
    buf.seek(0)
    return send_file(
        buf,
        as_attachment=True,
        download_name=f"pulse-appointments-{month}.xlsx",
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


db.connect()

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "8090")), debug=True)
